"""Convert a right-to-left Hebrew Markdown file (headings, blockquotes, one table, numbered items, **bold**, *italic*, `code`) to a DOCX.
Usage: python3 tools/md_he_to_docx.py docs/METHODOLOGY_section8_he.md exports/METHODOLOGY_section8_he.docx"""
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, dst = sys.argv[1], sys.argv[2]
lines = [l.rstrip('\n') for l in open(src, encoding='utf-8')]
lines = [l for l in lines if l.strip() not in ('<div dir="rtl" lang="he">', '</div>')]
doc = Document(); sec = doc.sections[0]; sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2)
FONT = 'Arial'
st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(11); st.element.rPr.rFonts.set(qn('w:cs'), FONT); st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

def rtl_par(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:bidi'); pPr.append(b); p.alignment = align
def run_rtl(r, size=None, bold=None, italic=None, mono=False, color=None):
    rPr = r._r.get_or_add_rPr(); x = OxmlElement('w:rtl'); rPr.append(x)
    r.font.name = FONT; rPr.rFonts.set(qn('w:cs'), FONT)
    if size: r.font.size = Pt(size); sz = OxmlElement('w:szCs'); sz.set(qn('w:val'), str(int(size * 2))); rPr.append(sz)
    if bold is not None: r.bold = bold; rPr.append(OxmlElement('w:bCs')) if bold else None
    if italic is not None: r.italic = italic
    if mono: r.font.name = 'Consolas'; rPr.rFonts.set(qn('w:cs'), 'Consolas')
    if color: r.font.color.rgb = RGBColor.from_string(color)
LRM = '\u200e'
RANGE = re.compile(r'(\d[\d.,:]*%?)\s?([–-])\s?(\d[\d.,:]*)')            # keep number ranges (1–5, 06:00–09:00, 0.02–0.05) left-to-right inside right-to-left text
def fix_ranges(t): return RANGE.sub(lambda m: f'{LRM}{m.group(1)}{LRM}{m.group(2)}{LRM}{m.group(3)}{LRM}', t)
TOK = re.compile(r'(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)')
def add_inline(p, text, size=None, bold_all=False, color=None):
    text = fix_ranges(text)
    for part in TOK.split(text):
        if not part: continue
        if part.startswith('**'): r = p.add_run(part[2:-2]); run_rtl(r, size, True, color=color)
        elif part.startswith('`'): r = p.add_run(part[1:-1]); run_rtl(r, size, bold_all or None, mono=True, color=color)
        elif part.startswith('*') and len(part) > 2: r = p.add_run(part[1:-1]); run_rtl(r, size, bold_all or None, True, color=color)
        else: r = p.add_run(part); run_rtl(r, size, bold_all or None, color=color)
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), hexcolor); tcPr.append(sh)
def para(text, size=11, space_after=6, align=WD_ALIGN_PARAGRAPH.RIGHT, indent=None, bold_all=False, color=None):
    p = doc.add_paragraph(); rtl_par(p, align); p.paragraph_format.space_after = Pt(space_after)
    if indent is not None: p.paragraph_format.right_indent = Cm(indent)
    add_inline(p, text, size, bold_all, color); return p

i = 0; quote = []; table = []
def flush_quote():
    global quote
    if not quote: return
    t = doc.add_table(rows=1, cols=1); t.alignment = 2; c = t.rows[0].cells[0]; shade(c, 'EAF1FB'); c.paragraphs[0]._p.getparent().remove(c.paragraphs[0]._p)
    for q in quote:
        if q.strip() == '': continue
        p = c.add_paragraph(); rtl_par(p); p.paragraph_format.space_after = Pt(3)
        txt = q[2:] if q.startswith('- ') else q
        add_inline(p, ('• ' if q.startswith('- ') else '') + txt, 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2); quote = []
def flush_table():
    global table
    if not table: return
    rows = [[c.strip() for c in r.strip().strip('|').split('|')] for r in table if not re.match(r'^\|[-| ]+\|$', r.strip())]
    t = doc.add_table(rows=len(rows), cols=2); t.style = 'Table Grid'; t.alignment = 2
    tblPr = t._tbl.tblPr; bv = OxmlElement('w:bidiVisual'); tblPr.append(bv)
    for ri, row in enumerate(rows):
        for ci, txt in enumerate(row):
            c = t.rows[ri].cells[ci]; c.width = Cm(5.2 if ci == 0 else 11.2); p = c.paragraphs[0]; rtl_par(p); p.paragraph_format.space_after = Pt(2)
            add_inline(p, txt, 10, bold_all=(ri == 0), color='FFFFFF' if ri == 0 else None)
            if ri == 0: shade(c, '1F3864')
    doc.add_paragraph().paragraph_format.space_after = Pt(4); table = []

for l in lines:
    if l.startswith('>'): flush_table(); quote.append(l[1:].lstrip() if l.strip() != '>' else ''); continue
    flush_quote()
    if l.startswith('|'): table.append(l); continue
    flush_table()
    if not l.strip() or l.strip() == '---': continue
    if l.startswith('# '): p = para(l[2:], 17, 10, bold_all=True, color='1F3864')
    elif l.startswith('## '): p = para(l[3:], 13.5, 6, bold_all=True, color='1F3864'); p.paragraph_format.space_before = Pt(12); p.paragraph_format.keep_with_next = True
    elif re.match(r'^\d+\. ', l):
        m = re.match(r'^(\d+)\. (.*)$', l); p = para(f'{m.group(1)}. ' + m.group(2), 10.5, 7)
    else: para(l, 11)
flush_quote(); flush_table()
doc.save(dst); print('saved', dst)
